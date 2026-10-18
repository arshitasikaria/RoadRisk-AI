"""PDF and Excel report generation."""
import io

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from . import analytics
from .constants import DATA_CREDIT, FEATURE_LABELS

NAVY = colors.HexColor("#1F4E79")
SEV_COLORS = {"Fatal": "#C0392B", "Serious": "#E67E22", "Minor": "#27AE60"}


def _condition_rows(prediction):
    rows = [["Road stretch (optional)", prediction.road_name or "-"]]
    rows += [[FEATURE_LABELS.get(k, k), v] for k, v in prediction.conditions.items()]
    return rows


def prediction_pdf(prediction):
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"RoadRisk AI Prediction #{prediction.pk}")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=NAVY, fontSize=20)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=NAVY)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10.5, leading=15)
    small = ParagraphStyle("s", parent=body, fontSize=8.5, textColor=colors.grey)
    big = ParagraphStyle("big", parent=ss["Title"], fontSize=26, textColor=colors.HexColor(SEV_COLORS.get(prediction.predicted_severity, "#000")))

    story = [Paragraph("RoadRisk AI - Accident Severity Report", h1),
             Paragraph(f"Report for prediction #{prediction.pk} &nbsp;|&nbsp; Generated {timezone.localtime().strftime('%d %b %Y, %H:%M')}"
                       f" &nbsp;|&nbsp; User: {prediction.user.username}", body), Spacer(1, 14), Paragraph("Input conditions", h2)]
    t = Table(_condition_rows(prediction), colWidths=[5.5 * cm, 10.5 * cm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#DCE9F5")),
                           ("FONTSIZE", (0, 0), (-1, -1), 10), ("PADDING", (0, 0), (-1, -1), 6)]))
    story += [t, Spacer(1, 14), Paragraph("Predicted severity", h2), Paragraph(prediction.predicted_severity.upper(), big),
              Paragraph(f"Confidence: {prediction.confidence * 100:.1f}% &nbsp;|&nbsp; Model used: {prediction.model_used}", body), Spacer(1, 8)]
    prob_rows = [["Severity", "Probability"]] + [[k, f"{v * 100:.1f}%"] for k, v in prediction.probabilities.items()]
    pt = Table(prob_rows, colWidths=[5 * cm, 4 * cm])
    pt.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("PADDING", (0, 0), (-1, -1), 5)]))
    story += [pt, Spacer(1, 14), Paragraph("AI road safety recommendations", h2)]
    for rec in prediction.recommendations.all():
        story.append(Paragraph("&bull; " + rec.suggestion_text, body))
    story += [Spacer(1, 18),
              Paragraph("<i>Predictions are probabilistic estimates based on historical highway accident data and should be reviewed "
                        "by road-safety experts before action is taken.</i>", small), Spacer(1, 4), Paragraph(DATA_CREDIT, small)]
    doc.build(story)
    return buf.getvalue()


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws):
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(45, max(len(str(c.value or "")) for c in col) + 3)


def prediction_excel(prediction):
    wb = Workbook()
    ws = wb.active
    ws.title = "Prediction"
    ws.append(["Field", "Value"])
    ws.append(["Prediction ID", prediction.pk])
    ws.append(["User", prediction.user.username])
    ws.append(["Created", timezone.localtime(prediction.created_at).strftime("%d %b %Y %H:%M")])
    for k, v in _condition_rows(prediction):
        ws.append([k, v])
    ws.append(["Predicted severity", prediction.predicted_severity])
    ws.append(["Confidence (%)", round(prediction.confidence * 100, 1)])
    ws.append(["Model used", prediction.model_used])
    _style_header(ws, 2)
    _autosize(ws)
    ws2 = wb.create_sheet("Probabilities")
    ws2.append(["Severity", "Probability (%)"])
    for k, v in prediction.probabilities.items():
        ws2.append([k, round(v * 100, 1)])
    _style_header(ws2, 2)
    ws3 = wb.create_sheet("Recommendations")
    ws3.append(["#", "Recommendation"])
    for i, r in enumerate(prediction.recommendations.all(), 1):
        ws3.append([i, r.suggestion_text])
    _style_header(ws3, 2)
    ws3.column_dimensions["B"].width = 120
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def analytics_excel(year=None, corridor=None):
    data = analytics.dashboard_data(year, corridor)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    k = data["kpis"]
    ws.append(["Metric", "Value"])
    for label, val in [("Corridor filter", corridor or "All corridors"), ("Year filter", year or "All years"),
                       ("Total accidents", k["total"]), ("Fatal", k["fatal"]), ("Serious", k["serious"]), ("Minor", k["minor"]),
                       ("Risk index (0-100)", k["risk_index"])]:
        ws.append([label, val])
    _style_header(ws, 2)
    _autosize(ws)

    def sheet(title, headers, rows):
        s = wb.create_sheet(title)
        s.append(headers)
        for r in rows:
            s.append(r)
        _style_header(s, len(headers))
        _autosize(s)

    cols = ("total", "fatal", "serious", "minor")
    sheet("Monthly Trend", ["Month", "Total", "Fatal", "Serious", "Minor"], [[m["month"], *[m[c] for c in cols]] for m in data["monthly"]])
    sheet("Yearly", ["Year", "Total", "Fatal", "Serious", "Minor"], [[y["year"], *[y[c] for c in cols]] for y in data["yearly"]])
    for key, title in (("weather", "Weather"), ("time_of_day", "Time of Day"), ("vehicle", "Vehicle"), ("cause", "Cause")):
        sheet(title, [title, "Total", "Fatal", "Serious", "Minor", "Risk index"],
              [[w["label"], *[w[c] for c in cols], w["risk_index"]] for w in data[key]])
    sheet("High-Risk Stretches", ["Road stretch", "Corridor", "Total", "Fatal", "Serious", "Minor", "Risk index", "Level"],
          [[r["road_name"], r["corridor"], r["total"], r["fatal"], r["serious"], r["minor"], r["risk_index"], r["risk_level"]] for r in data["roads"]])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
